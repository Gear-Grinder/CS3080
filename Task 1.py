import sys
import random
import requests
from bs4 import BeautifulSoup
import openpyxl
url = "https://www.ef.edu/english-resources/english-vocabulary/top-1000-words/"
response = requests.get(url, timeout=10)
soup = BeautifulSoup(response.text, "html.parser")
content_div = soup.find("div", class_="field-item even")
all_words = content_div.get_text().strip().split()
start_index = all_words.index("a")
common_1000 = all_words[start_index:start_index + 1000]
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = "Common Words"
sheet["A1"] = "Common Words"
for i, word in enumerate(common_1000, start=2):
    sheet[f"A{i}"] = word
excel_filename = "Common 1000 Words.xlsx"
workbook.save(excel_filename)
print(f"✅ Done! Saved {len(common_1000)} words to '{excel_filename}'.")
wb = openpyxl.load_workbook('Common 1000 Words.xlsx')
sheet = wb.active
words = []
for i in range(2, 1001):
    val = sheet.cell(row=i, column=1).value
    if val:
        words.append(val.strip())
word = words[random.randint(0, len(words) - 1)]
characters = list(word)
under = ["_"] * len(word)
tries = 5
letters = []
while tries > 0:
    print(" ".join(under))
    letter = input("Guess a letter: ").lower()
    if letter.isalpha() is False:
        continue
    if letter in letters:
        print("This letter was already chosen.\n")
        continue
    letters.append(letter)
    if letter in characters:
        for i, ch in enumerate(characters):
            if ch == letter:
                under[i] = letter
        print("\n" + " ".join(under))
    else:
        tries -= 1
        print(f"Wrong. Tries left: {tries}\n")
    if "_" not in under:
        print("You Won!")
        sys.exit()
print(f'\nAnswer: "{word}"')
print("You Lost!")
